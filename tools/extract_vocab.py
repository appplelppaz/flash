#!/usr/bin/env python3
"""Netflix / LLN subtitle file -> flashcard vocabulary candidates (TSV)."""
import argparse, os, re, sys
from collections import OrderedDict

ZH_STOP = set("""的 了 是 我 你 您 他 她 它 我们 你们 他们 她们 这 那 这个 那个 这些 那些 这样 那样 这里 那里
有 没 没有 在 和 跟 就 也 都 还 又 再 只 才 很 太 更 最 不 别 要 会 能 可以 可能 应该 得 着 过 把 被 让 给 对 从 向 到
吗 吧 呢 啊 呀 哦 嗯 哎 欸 唉 喂 嘛 咯 啦 呗 哈 嘿 什么 怎么 怎样 为什么 哪 哪里 哪个 谁 多少 几
因为 所以 但是 可是 而且 然后 如果 虽然 就是 还是 或者 而是 于是 不过
一 二 三 两 个 些 点 下 次 种 位 件 张 条 台 场 年 月 日 天 号
自己 别人 大家 一起 一下 一点 一样 一直 已经 现在 时候 今天 明天 昨天 上 下 里 中 外 前 后 左 右
好 好的 是的 对 不是 不用 没关系 谢谢 请 来 去 说 想 看 做 走 真 真的 当然 其实 反正 那么 这么 多 少 大 小
还有 是不是 有没有 不会 不能 只能 好不好 知不知道 怎么回事 干吗 干嘛 没到 有点 还好 不好意思 对不起
最好 样子 以后 之后 上面 下面 里面 外面 出来 回来 回去 过来 过去 起来 下去 这边 那边 那时候
需要 应该 觉得 知道 看到 看着 想要 开始 结束 帮 带 放 拿 找 叫 讲 说话 问 你好 请问 没事 没错""".split())
ZH_REDUP_OK = set("爸爸 妈妈 哥哥 姐姐 弟弟 妹妹 爷爷 奶奶 叔叔 阿姨 舅舅 姑姑 宝宝 娃娃 太太 星星 慢慢 刚刚 常常 天天 悄悄 偷偷 渐渐".split())
ZH_SURNAME = set("赵钱孙李周吴郑王冯陈卫蒋沈韩杨朱秦许何吕施张孔曹严华金魏陶姜谢邹章苏潘范彭鲁马苗方俞任袁柳唐罗薛雷贺汤殷郝安常乐傅齐康伍余元顾孟平黄穆萧尹姚邵汪毛米明成戴谈宋庞熊纪舒屈项祝董梁杜阮蓝季强贾路江童颜郭梅盛林钟徐邱骆高夏蔡田樊胡凌霍万柯管卢莫房宗丁邓单洪包左石崔龚程邢裴陆荣段富焦全秋仲宫宁甘厉武符刘景詹龙叶司黎薄白蒲赖卓屠蒙池乔谭")

EN_STOP = set("""a an the and or but so if then than that this these those i you he she it we they me him her us them my your his
its our their am is are was were be been being have has had do does did doing will would can could shall should may might must
not no nor of in on at to for with from by about into over after before between under above out up down off again once here
there when where why how all any both each few more most other some such only own same too very just now yeah yes ok okay oh
uh um hey hi hello well got get gets getting go goes going went gone come comes coming came know knows knew think thought like
likes want wants wanted see saw look looks looking say says said tell told make makes made take takes took give gives gave one
two three what who whom whose which don doesn didn isn aren wasn weren won gonna wanna gotta let lets right good bad thing
things people time""".split())
ES_STOP = set("""el la los las un una unos unas y o pero si no de del a al en con por para sin sobre entre desde hasta que
qué quien quién cual cuál como cómo cuando cuándo donde dónde porque yo tú él ella nosotros vosotros ellos ellas me te se nos
os le les lo mi mí tu ti su sus mis tus nuestro vuestro este esta esto estos estas ese esa eso esos esas aquel aquella ser soy
eres es somos sois son era eran fue fui fueron sido estar estoy estás está estamos están estaba estaban haber he has ha hemos
han había hay hacer hago haces hace hacen hizo tener tengo tienes tiene tienen tenía ir voy vas va vamos van poder puedo puedes
puede pueden querer quiero quieres quiere saber sabes sabe decir digo dices dice dijo ver veo ves muy más menos ya también
tampoco todo toda todos todas nada nadie algo alguien bien mal ahora aquí allí así sólo solo aunque gracias hola vale bueno
claro pues""".split())
FR_STOP = set("""le la les un une des du de au aux et ou mais si ne pas plus moins que qui quoi quand comment où parce pour
par avec sans sur sous dans chez vers depuis jusque je tu il elle on nous vous ils elles me te se lui leur mon ma mes ton ta
tes son sa ses notre nos votre vos leurs ce cet cette ces celui celle ceux être suis es est sommes êtes sont était étaient été
avoir ai as avons avez ont avait avaient eu faire fais fait faisons font aller vais vas va allons vont pouvoir peux peut
pouvons peuvent vouloir veux veut voulons veulent savoir sais sait savons savent dire dis dit disons disent voir vois voit
voyons voient très bien mal tout toute tous toutes rien personne quelque quelques déjà encore aussi alors donc ainsi ici là
oui non merci bonjour""".split())
STOP = {"zh": ZH_STOP, "en": EN_STOP, "es": ES_STOP, "fr": FR_STOP}

def read_rows(path):
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"): return _read_xlsx(path)
    if ext in (".csv", ".tsv", ".txt"): return _read_delim(path)
    if ext == ".srt": return _read_srt(path)
    raise SystemExit("unsupported file type: " + ext)

def _map_cols(hdr):
    idx = {"time": None, "text": None, "trans": None, "translit": None}
    for i, h in enumerate(hdr):
        if idx["time"] is None and h in ("time", "timestamp", "start"): idx["time"] = i
        elif idx["text"] is None and ("subtitle" in h or h in ("text", "line")): idx["text"] = i
        elif idx["trans"] is None and ("translation" in h or "translated" in h): idx["trans"] = i
        elif idx["translit"] is None and ("translit" in h or "pinyin" in h or "reading" in h): idx["translit"] = i
    return idx

def _row(idx, cells):
    def g(k):
        i = idx[k]
        return str(cells[i]).strip() if i is not None and i < len(cells) and cells[i] is not None else ""
    return {"time": g("time"), "text": g("text"), "trans": g("trans"), "translit": g("translit")}

def _read_xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["Subtitles"] if "Subtitles" in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    if not rows: return []
    idx = _map_cols([str(c or "").strip().lower() for c in rows[0]])
    body = rows[1:]
    if idx["text"] is None:
        idx, body = {"time": 0, "text": 1, "trans": 2, "translit": 3}, rows
    return [_row(idx, r) for r in body if r and any(r)]

def _read_delim(path):
    import csv
    with open(path, newline="", encoding="utf-8-sig") as f:
        sample = f.read(4096); f.seek(0)
        d = "\t" if sample.count("\t") > sample.count(",") else ","
        rows = [r for r in csv.reader(f, delimiter=d) if any(x.strip() for x in r)]
    if not rows: return []
    idx = _map_cols([c.strip().lower() for c in rows[0]])
    body = rows[1:]
    if idx["text"] is None:
        idx, body = {"time": 0, "text": 1, "trans": 2, "translit": 3}, rows
    return [_row(idx, r) for r in body]

def _read_srt(path):
    txt = open(path, encoding="utf-8-sig", errors="replace").read()
    out = []
    for block in re.split(r"\n\s*\n", txt):
        lines = [l for l in block.strip().splitlines() if l.strip()]
        if len(lines) < 2: continue
        tm, body = "", []
        for l in lines:
            if "-->" in l: tm = l.split("-->")[0].strip()
            elif re.fullmatch(r"\d+", l.strip()): continue
            else: body.append(l)
        if body: out.append({"time": tm, "text": " ".join(body).strip(), "trans": "", "translit": ""})
    return out

FR_HINT = set("le la les des du une est pas que qui vous nous je tu il elle ce cette pour dans avec plus mais on ne".split())
ES_HINT = set("el la los las un una es que de en por para con no se lo su muy pero como más está están qué".split())
EN_HINT = set("the and is are you i to of a in that it for on was with have this not be he she".split())

def detect_lang(rows):
    text = " ".join(r["text"] for r in rows[:600])
    cjk = sum(1 for ch in text if "一" <= ch <= "鿿")
    kana = sum(1 for ch in text if "぀" <= ch <= "ヿ")
    letters = sum(1 for ch in text if ch.isalpha())
    if cjk and cjk > kana and cjk / max(letters, 1) > 0.3: return "zh"
    ws = set(re.findall(r"[^\W\d_]+", text.lower(), re.UNICODE))
    score = {"fr": len(ws & FR_HINT) + 3 * text.count("œ") + text.count("ç") + sum(text.count(c) for c in "èêùôû"),
             "es": len(ws & ES_HINT) + 5 * (text.count("ñ") + text.count("¿") + text.count("¡")),
             "en": len(ws & EN_HINT)}
    score["fr"] += 2 * len(re.findall(r"\b[cdjlmnstqu]{1,2}['’]", text.lower()))
    return max(score, key=score.get)

def clean_line(s):
    s = s.replace("\n", " ").replace("\\N", " ")
    s = re.sub(r"[‎‏‪-‮﻿]", "", s)
    s = re.sub(r"<[^>]+>", "", s)
    s = re.sub(r"\[[^\]]*\]", " ", s)          # [man 1] [sighs]
    s = re.sub(r"\([^)]*\)", " ", s)           # (laughs)
    s = re.sub(r"^\s*[-–—]\s*", "", s)
    s = re.sub(r"^[A-Z][A-Za-z .'-]{0,18}:\s*", "", s)
    s = s.replace("♪", " ")
    return re.sub(r"\s+", " ", s).strip()

PARTICLES = set("up down out off on in over back away through around along across together apart forward".split())
BOUND_BAD = set("""a an the and or but so that which who whom whose this these those there here i you he she it we they me him
her us them my your his its our their am is are was were be been being have has had do does did doing will would can could
shall should may might must not no nor to of as if then than when where why how very just now yeah yes oh uh um well ok okay
el la los las un una y o que de del al en se lo le les mi tu su este esta ese esa es son era fue ser estar hay
le la les un une des du de et ou que qui ne pas je tu il elle on nous vous ils elles ce cet cette est sont était être avoir""".split())

def phrases_latin(raw, stop):
    """verb+particle pairs and clean recurring n-grams"""
    out = []
    for clause in re.split(r"[.!?,;:¡¿…\"“”]+", raw.replace("’", "'")):
        words = re.findall(r"[^\W\d_]+(?:'[^\W\d_]+)?", clause, re.UNICODE)
        if not words: continue
        low = [w.lower().strip("'") for w in words]
        propn = [i > 0 and w[:1].isupper() for i, w in enumerate(words)]
        for i in range(len(low) - 1):
            if (low[i + 1] in PARTICLES and low[i] not in stop and not propn[i]
                    and len(low[i]) > 2 and "'" not in low[i]):
                out.append((low[i] + " " + low[i + 1], True))
        for n in (2, 3, 4):
            for i in range(len(low) - n + 1):
                g = low[i:i + n]
                if any(propn[i:i + n]): continue
                if any("'" in t or len(t) < 2 for t in (g[0], g[-1])): continue
                if g[0] in BOUND_BAD or (g[-1] in BOUND_BAD and g[-1] not in PARTICLES): continue
                if all(t in stop for t in g): continue
                out.append((" ".join(g), False))
    return out

def phrases_zh(text):
    """jieba idiom tags + 4-character runs split into two 2-char tokens"""
    import jieba.posseg as pseg
    out = []
    for clause in re.split(r"[，。！？；：、…“”（）\s]+", text):
        toks = [(w, f) for w, f in pseg.cut(clause) if w.strip()]
        for w, f in toks:
            if f in ("i", "l") and len(w) >= 3: out.append((w, True))
        for i in range(len(toks) - 1):
            a, b = toks[i][0], toks[i + 1][0]
            if len(a) == 2 and len(b) == 2 and all("一" <= c <= "鿿" for c in a + b):
                out.append((a + b, False))
    return out

def tokens_zh(text):
    import jieba.posseg as pseg
    keep = ("n", "v", "a", "i", "l", "z", "s", "f", "b")
    drop = ("nr", "ns", "nt", "nw", "eng", "m", "q", "r", "p", "c", "u", "w", "x", "y", "e", "o", "t", "d")
    out = []
    for w, flag in pseg.cut(text):
        w = w.strip()
        if not w or flag in drop or flag.startswith(("nr", "ns", "nt")): continue
        if not all("一" <= ch <= "鿿" for ch in w): continue
        if not flag.startswith(keep): continue
        out.append(w)
    return out

def tokens_latin(text):
    text = text.lower().replace("’", "'")
    return [w.strip("'") for w in re.findall(r"[^\W\d_]+(?:'[^\W\d_]+)?", text, re.UNICODE) if len(w.strip("'")) >= 2]

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("infile")
    ap.add_argument("-o", "--out", default=None)
    ap.add_argument("--lang", default="auto", choices=["auto", "en", "zh", "es", "fr"])
    ap.add_argument("--max", type=int, default=0, help="0 = no limit (the intended setting)")
    ap.add_argument("--min-count", type=int, default=1)
    ap.add_argument("--order", default="appearance", choices=["appearance", "frequency"])
    ap.add_argument("--select", default="frequency", choices=["frequency", "first"])
    ap.add_argument("--min-len", type=int, default=0)
    ap.add_argument("--drop", default="")
    ap.add_argument("--keep-names", action="store_true")
    ap.add_argument("--phrases", action="store_true", default=True)
    ap.add_argument("--no-phrases", dest="phrases", action="store_false")
    ap.add_argument("--phrase-min-count", type=int, default=2)
    ap.add_argument("--example-max-chars", type=int, default=60)
    ap.add_argument("--min-example-chars", type=int, default=4)
    a = ap.parse_args()

    rows = [r for r in read_rows(a.infile) if clean_line(r["text"])]
    for r in rows:
        r["text"], r["trans"] = clean_line(r["text"]), clean_line(r["trans"])
    if not rows: raise SystemExit("no subtitle lines found")

    lang = a.lang if a.lang != "auto" else detect_lang(rows)
    dropset = set(x.strip() for x in a.drop.split(",") if x.strip())
    stop = STOP[lang]
    tok = tokens_zh if lang == "zh" else tokens_latin

    info = OrderedDict()
    for i, r in enumerate(rows):
        for w in tok(r["text"]):
            if w in stop or w in dropset: continue
            if lang == "zh" and not a.keep_names:
                if len(w) == 2 and w[0] == w[1] and w not in ZH_REDUP_OK: continue
                if len(w) == 1 and w in ZH_SURNAME: continue
            d = info.get(w)
            if d is None: info[w] = d = {"word": w, "count": 0, "first": i, "rows": []}
            d["count"] += 1
            if len(d["rows"]) < 8: d["rows"].append(i)

    ph = OrderedDict()
    if a.phrases:
        for i, r in enumerate(rows):
            for text_, always in (phrases_zh(r["text"]) if lang == "zh" else phrases_latin(r["text"], stop)):
                if text_ in dropset: continue
                d = ph.get(text_)
                if d is None: ph[text_] = d = {"word": text_, "count": 0, "first": i, "rows": [], "always": False}
                d["count"] += 1
                d["always"] = d["always"] or always
                if len(d["rows"]) < 8: d["rows"].append(i)

    def pick_example(d):
        best, best_score = None, -1e9
        for i in d["rows"]:
            r = rows[i]; n = len(r["text"])
            score = -abs(n - 20) if a.min_example_chars <= n <= a.example_max_chars else -abs(n - 20) - 40
            if r["trans"]: score += 50
            if i == d["first"]: score += 8
            if score > best_score: best, best_score = i, score
        return best if best is not None else d["first"]

    minlen = a.min_len or (2 if lang == "zh" else 3)
    items = [d for d in info.values()
             if d["count"] >= a.min_count and (len(d["word"]) >= minlen or d["count"] >= 2)]
    if a.max and len(items) > a.max:
        key = (lambda d: (-d["count"], d["first"])) if a.select == "frequency" else (lambda d: d["first"])
        items = sorted(items, key=key)[: a.max]
    items.sort(key=(lambda d: (-d["count"], d["first"])) if a.order == "frequency" else (lambda d: d["first"]))
    pitems = [d for d in ph.values() if d["always"] or d["count"] >= a.phrase_min_count]
    pitems.sort(key=lambda d: d["first"])

    out = a.out or (os.path.splitext(a.infile)[0] + "_%s_candidates.tsv" % lang)
    with open(out, "w", encoding="utf-8") as f:
        f.write("kind\tword\treading\tja\texample\texample_ja\tcount\tfirst_at\n")
        for kind, seq in (("word", items), ("phrase", pitems)):
            for d in seq:
                r = rows[pick_example(d)]
                f.write("\t".join([kind, d["word"], "", "", r["text"], r["trans"],
                                   str(d["count"]), r["time"]]) + "\n")
    sys.stderr.write("language=%s lines=%d words=%d phrases=%d -> %s\n"
                     % (lang, len(rows), len(items), len(pitems), out))
    print(lang)

if __name__ == "__main__":
    main()
